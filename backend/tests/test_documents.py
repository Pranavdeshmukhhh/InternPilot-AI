import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfReader

from app.documents import generate_cover_letter, generate_resume, parse_resume


@pytest.fixture
def candidate():
    return {
        "name": "Test Candidate",
        "email": "candidate@example.test",
        "phone": "+91 00000 00000",
        "country": "India",
        "degree": "B.Tech Computer Science",
        "university": "Test University",
        "graduation_year": 2029,
        "skills": ["HTML", "Python", "SQL"],
        "projects": [
            {"name": "Personal website", "description": "Built a static portfolio.", "skills": ["HTML"]},
            {"name": "Data notebook", "description": "Analyzed a public dataset using Python.", "skills": ["Python"]},
        ],
        "experience": [{"title": "Unverified Nobel Winner", "organization": "Invented", "description": "SECRET_UNVERIFIED_CLAIM"}],
        "verified_fields": ["name", "email", "country", "degree", "university", "graduation_year", "skills", "projects"],
    }


def pdf_text(content):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


def test_tailored_pdf_is_readable_and_uses_only_verified_existing_facts(candidate):
    result = generate_resume(candidate, {"title": "Python Intern", "skills": ["Python"], "description": "Python and SQL"})
    assert result["content"].startswith(b"%PDF-")
    text = pdf_text(result["content"])
    assert "Test Candidate" in text
    assert "Expected graduation: 2029" in text
    assert "Python, SQL, HTML" in text
    assert text.index("Data notebook") < text.index("Personal website")
    assert "SECRET_UNVERIFIED_CLAIM" not in text
    assert "+91 00000" not in text
    assert "experience" not in result["used_facts"]
    assert "phone" not in result["used_facts"]
    assert "NVIDIA" not in text


def test_imported_resume_is_unverified_and_export_roundtrip_never_auto_verifies(candidate):
    generated = generate_resume(candidate, {})
    imported = parse_resume(generated["content"], "resume.pdf")
    assert imported["verified"] is False
    assert imported["verified_fields"] == []
    assert imported["proposed_fields"]["email"] == candidate["email"]
    assert imported["proposed_fields"]["graduation_year"] == 2029
    assert "Python" in imported["proposed_fields"]["skills"]
    with pytest.raises(ValueError, match="Verify your name"):
        generate_resume(imported["proposed_fields"], {})


def test_txt_parser_requires_evidence_and_does_not_guess_graduation():
    parsed = parse_resume(b"Contact: a@example.test\nProject completed 2027. Python JavaScript\nhttps://github.com/example\nhttps://linkedin.com/in/example", "resume.txt")
    assert "graduation_year" not in parsed["proposed_fields"]
    assert parsed["proposed_fields"]["github"] == "https://github.com/example"
    assert parsed["proposed_fields"]["linkedin"] == "https://linkedin.com/in/example"
    assert "Java" not in parsed["proposed_fields"]["skills"]
    assert "JavaScript" in parsed["proposed_fields"]["skills"]
    assert parsed["evidence"]["email"] == "a@example.test"


def test_docx_reads_tables_without_treating_them_as_verified():
    doc = Document()
    doc.add_paragraph("Expected graduation: May 2029")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "cv@example.test"
    stream = io.BytesIO()
    doc.save(stream)
    result = parse_resume(stream.getvalue(), "resume.docx")
    assert result["proposed_fields"]["graduation_year"] == 2029
    assert result["proposed_fields"]["email"] == "cv@example.test"
    assert "Python" in result["proposed_fields"]["skills"]
    assert result["verified_fields"] == []


@pytest.mark.parametrize("filename,data", [("resume.exe", b"foo"), ("resume.pdf", b"not pdf"), ("resume.docx", b"not zip"), ("resume.txt", b"\xff\xfe"), ("resume.txt", b""), ("resume.txt", b"x" * (8 * 1024 * 1024 + 1))], ids=["unsupported-extension", "invalid-pdf", "invalid-docx", "invalid-encoding", "empty-file", "oversized-file"])
def test_invalid_input_is_rejected(filename, data):
    with pytest.raises(ValueError):
        parse_resume(data, filename)


def test_macro_docx_is_rejected():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/vbaProject.bin", "macro")
    with pytest.raises(ValueError, match="Macro"):
        parse_resume(stream.getvalue(), "resume.docx")


def test_missing_verified_contact_and_malformed_contact_fail_closed(candidate):
    candidate["verified_fields"].remove("email")
    with pytest.raises(ValueError, match="Verify a valid contact email"):
        generate_resume(candidate, {})
    candidate["verified_fields"].append("email")
    candidate["email"] = "not-an-email"
    with pytest.raises(ValueError, match="Verify a valid contact email"):
        generate_cover_letter(candidate, {})


def test_resume_escapes_markup_instead_of_executing_it(candidate):
    candidate["projects"] = [{"name": "<b>Literal project</b>", "description": "5 < 7 & 9 > 2"}]
    generated = generate_resume(candidate, {})
    text = pdf_text(generated["content"])
    assert "<b>Literal project</b>" in text
    assert "5 < 7 & 9 > 2" in text


def test_cover_letter_contains_context_and_only_verified_candidate_facts(candidate):
    candidate["verified_fields"].remove("skills")
    candidate["verified_fields"].remove("projects")
    candidate["skills"] = ["SECRET_UNVERIFIED_SKILL"]
    letter = generate_cover_letter(candidate, {"title": "ML Research Intern", "company": "Example Research"})
    text = pdf_text(letter["content"])
    assert "ML Research Intern at Example Research" in text
    assert "expected graduation year is 2029" in " ".join(text.split())
    assert "SECRET_UNVERIFIED" not in text
    assert "Data notebook" not in text
    assert "expert" not in text.lower()
    assert "experience" not in letter["used_facts"]
    assert "skills" not in letter["used_facts"]


def test_no_education_claims_from_defaults_without_verification(candidate):
    candidate["verified_fields"] = ["name", "email"]
    result = generate_resume(candidate, {})
    assert "2029" not in result["text"]
    assert "India" not in result["text"]
    assert "B.Tech" not in result["text"]
    assert result["used_facts"] == ["name", "email"]


def test_embedded_font_preserves_supported_unicode_names(candidate):
    candidate["name"] = "Renée Jørgensen"
    result = generate_resume(candidate, {})
    assert "Renée Jørgensen" in pdf_text(result["content"])
